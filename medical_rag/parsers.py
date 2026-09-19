# -*- coding: utf-8 -*-
"""文档解析层：把 PDF / DOCX / Markdown / TXT / QA-JSONL 统一成 :class:`RawDoc`。

本模块是整条医疗 RAG 链路的入口，只负责一件事：**把原始文件拆成有序的 block，
并且尽最大努力保住出处信息**（页码、标题层级、是不是表格）。切块、索引、重排
都在后面的模块里做，这里不做任何语义处理。

两个设计决定值得先说清楚：

1. PDF 的表格必须单独抽。医疗文档里最值钱的往往是剂量表、分级表、禁忌表，
   而 PDF 的表格在按行取文本时会被拍扁成一串没有列边界的碎字符，接到下游就是
   不可用的噪声。所以这里用 pymupdf 逐页取正文、用 pdfplumber 单独抽表格，
   表格转成 Markdown 后作为 ``content_type="table"`` 的独立 block。

2. 页码必须逐页记。医疗回答如果不能回溯到"第几页"，引用就没有可核查性。
   因此正文 block 不跨页合并——宁可把一个跨页段落切成两块，也不让页码变模糊。

所有第三方依赖（pymupdf / pdfplumber / python-docx）都是可选的。缺任何一个都
只降级、打 warning，绝不让 ``import medical_rag.parsers`` 失败——开发机往往只装
了其中一两个。

用法::

    python medical_rag/parsers.py --input data/guideline.pdf \\
        --source_tier L1_guideline --output out/raw_docs.jsonl
"""

import argparse
import json
import logging
import os
import re
import sys
from dataclasses import asdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:  # 兼容两种运行方式：作为包被 import，或直接 python medical_rag/parsers.py
    from medical_rag.schema import (
        ALL_TIERS,
        CONTENT_QA,
        CONTENT_TABLE,
        CONTENT_TEXT,
        RawDoc,
        normalize_text,
        read_jsonl,
        stable_id,
        write_jsonl,
    )
except ImportError:  # pragma: no cover - 直接执行脚本时补上仓库根目录
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from medical_rag.schema import (
        ALL_TIERS,
        CONTENT_QA,
        CONTENT_TABLE,
        CONTENT_TEXT,
        RawDoc,
        normalize_text,
        read_jsonl,
        stable_id,
        write_jsonl,
    )

logger = logging.getLogger(__name__)


# --------------------------------------------------------------------------
# 可选依赖：只探测，不在 import 期报警
# --------------------------------------------------------------------------
# 警告延迟到"真的需要它却没有"的那一刻再打。整个 medical_rag 有六七个模块，
# 如果每个模块都在 import 期把缺失的依赖吼一遍，日志会被淹没。

try:  # pymupdf 新旧包名都兼容：新版推荐 import pymupdf，老版只有 fitz
    import pymupdf as _fitz
except ImportError:
    try:
        import fitz as _fitz
    except ImportError:
        _fitz = None

try:
    import pdfplumber as _pdfplumber
except ImportError:
    _pdfplumber = None

try:
    import docx as _docx
except ImportError:
    _docx = None


_WARNED = set()


def _warn_once(key: str, message: str) -> None:
    """同一类降级只提示一次，避免逐文件刷屏。"""
    if key in _WARNED:
        return
    _WARNED.add(key)
    logger.warning(message)


# --------------------------------------------------------------------------
# 标题启发式
# --------------------------------------------------------------------------
# 标题只有一行、没有句末标点、经常带编号。下面三条规则按"信号强度"排序使用。

HEADING_MAX_CHARS = 40

# 带编号的行是最强的标题信号：第3章 / 第三节 / 1.2.3 / （一） / 一、 / IV.
_NUM_HEADING_RE = re.compile(
    r"^("
    r"第\s*[0-9一二三四五六七八九十百千零两]+\s*[章节篇部条款讲编]"
    r"|[0-9]+(?:\.[0-9]+){0,4}\s*[\.、．\s]"
    r"|[（(]\s*[0-9一二三四五六七八九十]+\s*[)）]"
    r"|[一二三四五六七八九十]+\s*[、.．]"
    r"|[IVXivx]{1,5}\s*[、.．]"
    r")"
)

# 句末标点出现在行内就说明这是正文而不是标题
_SENT_END_RE = re.compile(r"[。！？；!?;]")

# 英文正文以 . 收尾。句点只在行尾才算句末标点，否则 "1.2 口服降糖药"
# 这类点分编号标题会被整片误判成正文。
_ASCII_SENT_END_RE = re.compile(r"\.\s*$")

# 页眉页脚里的孤立页码：12 / - 12 - / 第 12 页 / 3/58
_PAGE_NOISE_RE = re.compile(
    r"^[\s\-—–_·.]*(?:第?\s*[0-9]{1,4}\s*页?|[0-9]{1,4}\s*/\s*[0-9]{1,4})[\s\-—–_·.]*$"
)

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s*(\S.*)$")
_MD_HR_RE = re.compile(r"^(?:-{3,}|\*{3,}|_{3,})$")
_DOCX_HEADING_RE = re.compile(r"(?:heading|标题)\s*([1-6])", re.IGNORECASE)

_TEXT_SUFFIXES = {".txt", ".text", ".log"}
_MD_SUFFIXES = {".md", ".markdown", ".mdown"}
_QA_SUFFIXES = {".jsonl", ".json"}
SUPPORTED_SUFFIXES = (
    {".pdf", ".docx"} | _TEXT_SUFFIXES | _MD_SUFFIXES | _QA_SUFFIXES
)


def _is_cjk(char: str) -> bool:
    """是否属于 CJK 区（含中日韩统一表意文字、中文标点、全角符号）。

    用码位区间而不是字面量，是为了让边界一眼可读、也不受源文件编码影响。
    """
    code = ord(char)
    return (
        0x4E00 <= code <= 0x9FFF  # CJK 统一表意文字
        or 0x3000 <= code <= 0x303F  # 中文标点
        or 0xFF00 <= code <= 0xFFEF  # 全角字符
    )


def _heading_level_from_prefix(text: str) -> int:
    """从编号前缀推层级：章 > 节 > 条，点分编号按点数定深度。"""
    if re.match(r"^第\s*[0-9一二三四五六七八九十百千零两]+\s*[章篇编]", text):
        return 1
    if re.match(r"^第\s*[0-9一二三四五六七八九十百千零两]+\s*[节部讲]", text):
        return 2
    if re.match(r"^第\s*[0-9一二三四五六七八九十百千零两]+\s*[条款]", text):
        return 3
    dotted = re.match(r"^([0-9]+(?:\.[0-9]+)*)", text)
    if dotted:
        return min(6, dotted.group(1).count(".") + 1)
    if re.match(r"^[一二三四五六七八九十]+\s*[、.．]", text):
        return 2
    if re.match(r"^[（(]", text):
        return 3
    return 1


def _looks_like_heading(text: str, width_hint: Optional[int] = None) -> Tuple[bool, int]:
    """判断一行是不是标题，返回 (是否标题, 层级)。

    ``width_hint`` 是本页正文的典型行宽。PDF 的正文被硬换行切成很多短行，这些
    短行同样没有句末标点，光靠"无标点"会把整段正文全判成标题。正文行一般接近
    满行宽，标题行明显更短，所以在有行宽参考时额外要求"明显短于正文"。
    """
    text = normalize_text(text)
    if not text or len(text) >= HEADING_MAX_CHARS:
        return False, 0
    if _NUM_HEADING_RE.match(text):
        return True, _heading_level_from_prefix(text)
    if len(text) < 2 or _SENT_END_RE.search(text) or _ASCII_SENT_END_RE.search(text):
        return False, 0
    if text[-1] in "，,、:：的和与及或":
        return False, 0
    if width_hint and len(text) > max(0.7 * width_hint, 12):
        return False, 0
    return True, 1


def _make_block(
    text: str,
    page: Optional[int] = None,
    is_heading: bool = False,
    heading_level: int = 0,
    content_type: str = CONTENT_TEXT,
) -> Dict[str, Any]:
    """统一的 block 构造口，保证下游拿到的键完全一致。"""
    return {
        "text": text,
        "page": page,
        "is_heading": bool(is_heading),
        "heading_level": int(heading_level),
        "content_type": content_type,
    }


def _join_lines(lines: List[str]) -> str:
    """把被硬换行切开的同一段正文拼回去。

    中文换行处不该补空格（补了会在词中间插空格），英文换行处必须补空格
    （不补会把两个单词粘成一个），所以按两端字符类型决定。
    """
    out = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if not out:
            out = line
            continue
        if _is_cjk(out[-1]) or _is_cjk(line[0]):
            out += line
        else:
            out += " " + line
    return out


def _blocks_from_lines(lines: List[str], page: Optional[int] = None) -> List[Dict[str, Any]]:
    """把一页/一个文件的裸文本行切成标题块和段落块。"""
    kept = [ln.strip() for ln in lines]
    kept = [ln for ln in kept if ln and not _PAGE_NOISE_RE.match(ln)]
    if not kept:
        return []

    # 用中位行宽做参考，避免被个别超长行带偏
    widths = sorted(len(ln) for ln in kept)
    width_hint = widths[len(widths) // 2] if len(widths) >= 4 else None

    blocks = []  # type: List[Dict[str, Any]]
    buffer = []  # type: List[str]

    def flush():
        if not buffer:
            return
        text = normalize_text(_join_lines(buffer))
        del buffer[:]
        if text:
            blocks.append(_make_block(text, page=page))

    for line in kept:
        is_heading, level = _looks_like_heading(line, width_hint=width_hint)
        if is_heading:
            flush()
            blocks.append(
                _make_block(normalize_text(line), page=page, is_heading=True, heading_level=level)
            )
        else:
            buffer.append(line)
    flush()
    return blocks


# --------------------------------------------------------------------------
# 表格 -> Markdown
# --------------------------------------------------------------------------
def _clean_cell(value: Any) -> str:
    """单元格内的换行必须压掉，否则会破坏 Markdown 表格的行结构。"""
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    text = text.replace("|", "\\|")
    return normalize_text(text)


def table_to_markdown(rows: Iterable[Iterable[Any]]) -> str:
    """把二维表转成 Markdown 表格文本；空表返回空串。

    保留 Markdown 形态而不是拍成纯文本，是因为列对应关系就是剂量表的全部信息：
    "二甲双胍 | 起始 500mg | 最大 2000mg" 一旦丢掉分隔符就无法还原。
    """
    cleaned = []  # type: List[List[str]]
    for row in rows or []:
        if row is None:
            continue
        cells = [_clean_cell(cell) for cell in row]
        if not any(cells):
            continue
        cleaned.append(cells)
    if not cleaned:
        return ""

    width = max(len(row) for row in cleaned)
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]

    # 全空行在上面已被滤掉，所以首行必定至少有一个非空单元格，可以直接当表头
    header = cleaned[0]
    body = cleaned[1:]

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _pdf_text_blocks(path: str) -> List[List[Dict[str, Any]]]:
    """逐页取正文，返回 ``[第1页的blocks, 第2页的blocks, ...]``。"""
    if _fitz is not None:
        try:
            document = _fitz.open(path)
        except Exception as exc:
            logger.warning("pymupdf 打开失败，跳过正文抽取 %s: %s", path, exc)
            return []
        pages = []  # type: List[List[Dict[str, Any]]]
        try:
            for page_no, page in enumerate(document, start=1):
                try:
                    raw = page.get_text("text") or ""
                except Exception as exc:
                    logger.warning("%s 第 %d 页正文抽取失败: %s", path, page_no, exc)
                    raw = ""
                pages.append(_blocks_from_lines(raw.splitlines(), page=page_no))
        finally:
            try:
                document.close()
            except Exception:
                pass
        return pages

    if _pdfplumber is not None:
        _warn_once("pdf_no_fitz", "未安装 pymupdf，PDF 正文改用 pdfplumber 抽取（速度更慢）")
        pages = []
        try:
            with _pdfplumber.open(path) as pdf:
                for page_no, page in enumerate(pdf.pages, start=1):
                    try:
                        raw = page.extract_text() or ""
                    except Exception as exc:
                        logger.warning("%s 第 %d 页正文抽取失败: %s", path, page_no, exc)
                        raw = ""
                    pages.append(_blocks_from_lines(raw.splitlines(), page=page_no))
        except Exception as exc:
            logger.warning("pdfplumber 打开失败，跳过正文抽取 %s: %s", path, exc)
            return []
        return pages

    _warn_once("pdf_none", "pymupdf 与 pdfplumber 均不可用，PDF 正文无法抽取，已跳过")
    return []


def _pdf_table_blocks(path: str) -> Dict[int, List[Dict[str, Any]]]:
    """逐页抽表格，返回 ``{页码: [表格block, ...]}``。"""
    if _pdfplumber is None:
        _warn_once("pdf_no_plumber", "未安装 pdfplumber，PDF 表格不再单独抽取（剂量表可能被拍扁）")
        return {}

    tables_by_page = {}  # type: Dict[int, List[Dict[str, Any]]]
    try:
        with _pdfplumber.open(path) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                try:
                    raw_tables = page.extract_tables() or []
                except Exception as exc:
                    # 单页表格解析失败不应该毁掉整本文档
                    logger.warning("%s 第 %d 页表格抽取失败: %s", path, page_no, exc)
                    continue
                blocks = []
                for raw_table in raw_tables:
                    markdown = table_to_markdown(raw_table)
                    if markdown:
                        blocks.append(
                            _make_block(markdown, page=page_no, content_type=CONTENT_TABLE)
                        )
                if blocks:
                    tables_by_page[page_no] = blocks
    except Exception as exc:
        logger.warning("pdfplumber 打开失败，跳过表格抽取 %s: %s", path, exc)
        return {}
    return tables_by_page


def parse_pdf(path: str) -> List[Dict[str, Any]]:
    """解析 PDF：正文按页顺序排列，每页正文后面紧跟该页的表格块。"""
    text_pages = _pdf_text_blocks(path)
    tables_by_page = _pdf_table_blocks(path)

    total_pages = max(len(text_pages), max(tables_by_page.keys()) if tables_by_page else 0)
    blocks = []  # type: List[Dict[str, Any]]
    for page_no in range(1, total_pages + 1):
        if page_no <= len(text_pages):
            blocks.extend(text_pages[page_no - 1])
        blocks.extend(tables_by_page.get(page_no, []))
    return blocks


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def _docx_heading_level(style_name: str) -> int:
    """从样式名推标题层级；识别不出来返回 0（当正文处理）。"""
    if not style_name:
        return 0
    match = _DOCX_HEADING_RE.search(style_name)
    if match:
        return int(match.group(1))
    lowered = style_name.strip().lower()
    if lowered in ("title", "subtitle") or style_name.strip() in ("标题", "副标题"):
        return 1
    return 0


def _iter_docx_body(document) -> List[Tuple[str, Any]]:
    """按正文 XML 顺序返回 (类型, 对象)，保证表格出现在原来的位置。

    python-docx 的 ``paragraphs`` 和 ``tables`` 是两个互不相干的列表，直接先段后表
    会把剂量表全部甩到文末，和上文脱节。所以直接遍历 body 的子元素。
    """
    try:
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        items = []  # type: List[Tuple[str, Any]]
        for child in document.element.body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                items.append(("p", Paragraph(child, document)))
            elif tag == "tbl":
                items.append(("tbl", Table(child, document)))
        return items
    except Exception as exc:
        # 兜底：顺序会丢，但总比整篇解析失败强
        _warn_once("docx_order", "无法按原顺序遍历 DOCX 正文，降级为先段落后表格: %s" % exc)
        items = [("p", para) for para in document.paragraphs]
        items.extend(("tbl", table) for table in document.tables)
        return items


def parse_docx(path: str) -> List[Dict[str, Any]]:
    """解析 DOCX：用 paragraph.style.name 里的 Heading N 定标题层级。"""
    if _docx is None:
        _warn_once("docx_none", "未安装 python-docx，DOCX 文件无法解析，已跳过")
        return []
    try:
        document = _docx.Document(path)
    except Exception as exc:
        logger.warning("python-docx 打开失败，跳过 %s: %s", path, exc)
        return []

    blocks = []  # type: List[Dict[str, Any]]
    for kind, item in _iter_docx_body(document):
        try:
            if kind == "p":
                text = normalize_text(item.text)
                if not text:
                    continue
                style_name = ""
                try:
                    style_name = item.style.name or ""
                except Exception:
                    style_name = ""
                level = _docx_heading_level(style_name)
                if level:
                    blocks.append(_make_block(text, is_heading=True, heading_level=level))
                    continue
                # 样式没标成标题时再用文本启发式兜一层
                is_heading, guessed = _looks_like_heading(text)
                blocks.append(
                    _make_block(text, is_heading=is_heading, heading_level=guessed if is_heading else 0)
                )
            else:
                rows = [[cell.text for cell in row.cells] for row in item.rows]
                markdown = table_to_markdown(rows)
                if markdown:
                    blocks.append(_make_block(markdown, content_type=CONTENT_TABLE))
        except Exception as exc:
            logger.warning("DOCX 元素解析失败，已跳过该元素 %s: %s", path, exc)
    return blocks


# --------------------------------------------------------------------------
# Markdown / TXT
# --------------------------------------------------------------------------
def parse_markdown(text: str) -> List[Dict[str, Any]]:
    """纯正则解析 Markdown，不依赖 markdown / bs4（它们大概率没装）。"""
    blocks = []  # type: List[Dict[str, Any]]
    lines = text.splitlines()
    para = []  # type: List[str]

    def flush():
        if not para:
            return
        merged = normalize_text(_join_lines(para))
        del para[:]
        if merged:
            blocks.append(_make_block(merged))

    idx = 0
    total = len(lines)
    while idx < total:
        raw = lines[idx]
        stripped = raw.strip()

        # 代码块整体作为一个文本块，块内的 # 不是标题
        if stripped.startswith("```") or stripped.startswith("~~~"):
            flush()
            fence = stripped[:3]
            idx += 1
            code = []
            while idx < total and not lines[idx].strip().startswith(fence):
                code.append(lines[idx])
                idx += 1
            idx += 1
            body = "\n".join(code).strip()
            if body:
                blocks.append(_make_block(body))
            continue

        if not stripped:
            flush()
            idx += 1
            continue

        heading = _MD_HEADING_RE.match(stripped)
        if heading:
            flush()
            title = normalize_text(heading.group(2).strip().strip("#").strip())
            if title:
                blocks.append(_make_block(title, is_heading=True, heading_level=len(heading.group(1))))
            idx += 1
            continue

        # setext 标题：正文行下面单独一行 === 或 ---
        if para and re.match(r"^=+$", stripped):
            title = normalize_text(_join_lines(para))
            del para[:]
            if title:
                blocks.append(_make_block(title, is_heading=True, heading_level=1))
            idx += 1
            continue
        if len(para) == 1 and re.match(r"^-{2,}$", stripped):
            title = normalize_text(para[0])
            del para[:]
            if title:
                blocks.append(_make_block(title, is_heading=True, heading_level=2))
            idx += 1
            continue

        if not para and _MD_HR_RE.match(stripped):
            idx += 1
            continue

        # 连续的 | 开头行就是一张表，原样保留 Markdown 形态
        if stripped.startswith("|"):
            flush()
            table_lines = []
            while idx < total and lines[idx].strip().startswith("|"):
                table_lines.append(normalize_text(lines[idx].strip()))
                idx += 1
            body = "\n".join(ln for ln in table_lines if ln)
            if body:
                blocks.append(_make_block(body, content_type=CONTENT_TABLE))
            continue

        para.append(stripped)
        idx += 1

    flush()
    return blocks


def parse_txt(text: str) -> List[Dict[str, Any]]:
    """TXT 按空行分段；单行短段落再走一次标题启发式。"""
    blocks = []  # type: List[Dict[str, Any]]
    for raw_para in re.split(r"\n\s*\n", text):
        lines = [ln.strip() for ln in raw_para.splitlines()]
        lines = [ln for ln in lines if ln and not _PAGE_NOISE_RE.match(ln)]
        if not lines:
            continue
        if len(lines) == 1:
            # 独占一段的短行是很强的标题信号，此时不需要行宽参考
            is_heading, level = _looks_like_heading(lines[0])
            text_one = normalize_text(lines[0])
            if text_one:
                blocks.append(
                    _make_block(
                        text_one,
                        is_heading=is_heading,
                        heading_level=level if is_heading else 0,
                    )
                )
            continue
        blocks.extend(_blocks_from_lines(lines))
    return blocks


def read_text_file(path: str) -> str:
    """读纯文本，按 Windows 中文环境常见编码依次尝试。"""
    candidates = ("utf-8-sig", "utf-8", "gb18030", "utf-16")
    for encoding in candidates:
        try:
            return Path(path).read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError, LookupError):
            continue
    # 全都不行就强读，宁可丢几个字也不要整篇丢掉
    return Path(path).read_text(encoding="utf-8", errors="ignore")


# --------------------------------------------------------------------------
# 对外入口
# --------------------------------------------------------------------------
def load_document(
    path: str,
    source_tier: str,
    doc_title: Optional[str] = None,
    version: str = "",
) -> RawDoc:
    """解析单个文档为 :class:`RawDoc`。

    永远返回 RawDoc，绝不抛异常：文件损坏时返回 ``blocks == []`` 并在
    ``extra["parse_error"]`` 里记下原因，调用方按空 blocks 跳过即可。这样批量
    入库时一个坏文件不会中断整批。
    """
    path_obj = Path(path)
    title = normalize_text(doc_title) if doc_title else path_obj.stem
    # doc_id 只由路径/标题/版本决定，重复入库时 id 不会漂移
    doc_id = stable_id("doc", str(path_obj), title, version)
    doc = RawDoc(
        doc_id=doc_id,
        doc_title=title,
        source_tier=source_tier,
        source_path=str(path_obj),
        version=version,
        blocks=[],
        extra={},
    )

    suffix = path_obj.suffix.lower()
    try:
        if not path_obj.is_file():
            raise IOError("文件不存在或不是普通文件")
        if suffix == ".pdf":
            doc.blocks = parse_pdf(str(path_obj))
        elif suffix == ".docx":
            doc.blocks = parse_docx(str(path_obj))
        elif suffix in _MD_SUFFIXES:
            doc.blocks = parse_markdown(read_text_file(str(path_obj)))
        elif suffix == ".doc":
            raise IOError("不支持旧版 .doc，请先另存为 .docx")
        else:
            # 未知后缀一律当纯文本读，读不出来会在下面被捕获
            doc.blocks = parse_txt(read_text_file(str(path_obj)))
    except Exception as exc:
        logger.warning("解析失败，已跳过 %s: %s", path_obj, exc)
        doc.extra["parse_error"] = str(exc)
        doc.blocks = []
        return doc

    doc.extra["n_blocks"] = len(doc.blocks)
    doc.extra["n_tables"] = sum(1 for b in doc.blocks if b.get("content_type") == CONTENT_TABLE)
    pages = [b.get("page") for b in doc.blocks if b.get("page")]
    doc.extra["n_pages"] = max(pages) if pages else 0
    if not doc.blocks:
        logger.warning("未解析出任何内容: %s", path_obj)
    return doc


def load_qa_jsonl(
    path: str,
    source_tier: str,
    question_field: str = "question",
    answer_field: str = "answer",
    title_field: Optional[str] = None,
    max_rows: int = -1,
) -> List[RawDoc]:
    """把 QA 语料（每行一条 ``{"question":..., "answer":...}``）转成 RawDoc 列表。

    一条 QA 就是一个独立的 doc：问答之间没有篇章结构，强行合并成一篇只会让
    切块把不相干的两条问答粘在一起。问答拼成一个 block 是因为只留答案会丢掉
    "在问什么"这个检索锚点。
    """
    docs = []  # type: List[RawDoc]
    try:
        rows = read_jsonl(path, max_rows=max_rows)
    except Exception as exc:
        logger.warning("QA 文件打开失败，已跳过 %s: %s", path, exc)
        return docs

    index = 0
    while True:
        # 逐行推进并单独兜异常：某一行 JSON 坏掉不应该让整个文件报废
        try:
            row = next(rows)
        except StopIteration:
            break
        except Exception as exc:
            logger.warning("QA 文件解析中断 %s: %s", path, exc)
            break

        index += 1
        if not isinstance(row, dict):
            continue
        question = normalize_text(row.get(question_field, ""))
        answer = normalize_text(row.get(answer_field, ""))
        if not question or not answer:
            logger.warning("第 %d 行缺少 %s/%s 字段，已跳过: %s", index, question_field, answer_field, path)
            continue

        title = ""
        if title_field:
            title = normalize_text(row.get(title_field, ""))
        if not title:
            title = question[:40]

        doc_id = stable_id("qa", path, index, question)
        text = "问：%s\n答：%s" % (question, answer)
        docs.append(
            RawDoc(
                doc_id=doc_id,
                doc_title=title,
                source_tier=source_tier,
                source_path=str(path),
                version="",
                blocks=[
                    _make_block(text, page=None, is_heading=False, heading_level=0, content_type=CONTENT_QA)
                ],
                extra={"row_index": index},
            )
        )
    return docs


def iter_input_files(input_path: str, recursive: bool = True) -> List[Path]:
    """输入是文件就返回它自己；是目录就按支持的后缀收集。"""
    target = Path(input_path)
    if target.is_file():
        return [target]
    if not target.is_dir():
        logger.warning("输入路径不存在: %s", input_path)
        return []
    pattern = "**/*" if recursive else "*"
    files = [p for p in sorted(target.glob(pattern)) if p.is_file()]
    return [p for p in files if p.suffix.lower() in SUPPORTED_SUFFIXES]


def load_path(
    input_path: str,
    source_tier: str,
    version: str = "",
    recursive: bool = True,
    question_field: str = "question",
    answer_field: str = "answer",
    title_field: Optional[str] = None,
    max_rows: int = -1,
) -> List[RawDoc]:
    """批量入口：文件或目录 -> RawDoc 列表，jsonl/json 自动走 QA 分支。"""
    docs = []  # type: List[RawDoc]
    for file_path in iter_input_files(input_path, recursive=recursive):
        if file_path.suffix.lower() in _QA_SUFFIXES:
            docs.extend(
                load_qa_jsonl(
                    str(file_path),
                    source_tier,
                    question_field=question_field,
                    answer_field=answer_field,
                    title_field=title_field,
                    max_rows=max_rows,
                )
            )
        else:
            doc = load_document(str(file_path), source_tier, version=version)
            if doc.blocks:
                docs.append(doc)
    return docs


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def parse_args():
    parser = argparse.ArgumentParser(
        description="Parse medical documents (PDF/DOCX/MD/TXT/QA-JSONL) into RawDoc JSONL."
    )
    parser.add_argument("--input", required=True, help="Input file or directory.")
    parser.add_argument(
        "--source_tier",
        required=True,
        choices=list(ALL_TIERS),
        help="Source tier of the corpus; drives the authority weight downstream.",
    )
    parser.add_argument("--output", required=True, help="Output RawDoc JSONL path.")
    parser.add_argument("--version", default="", help="Document version, e.g. 2023.")
    parser.add_argument("--no_recursive", action="store_true", help="Do not walk sub-directories.")
    parser.add_argument("--question_field", default="question", help="QA JSONL question field.")
    parser.add_argument("--answer_field", default="answer", help="QA JSONL answer field.")
    parser.add_argument("--title_field", default=None, help="QA JSONL title field (optional).")
    parser.add_argument("--max_rows", type=int, default=-1, help="Max QA rows per file; -1 for all.")
    parser.add_argument("--print_samples", type=int, default=3, help="Number of blocks to preview.")
    return parser.parse_args()


def main():
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    args = parse_args()

    docs = load_path(
        args.input,
        args.source_tier,
        version=args.version,
        recursive=not args.no_recursive,
        question_field=args.question_field,
        answer_field=args.answer_field,
        title_field=args.title_field,
        max_rows=args.max_rows,
    )
    written = write_jsonl(args.output, (asdict(doc) for doc in docs))

    stats = {
        "input": args.input,
        "output": args.output,
        "source_tier": args.source_tier,
        "n_docs": written,
        "n_blocks": sum(len(doc.blocks) for doc in docs),
        "n_heading_blocks": sum(1 for doc in docs for b in doc.blocks if b.get("is_heading")),
        "n_table_blocks": sum(
            1 for doc in docs for b in doc.blocks if b.get("content_type") == CONTENT_TABLE
        ),
    }
    print(json.dumps(stats, ensure_ascii=False, indent=2))

    if docs and args.print_samples > 0:
        print("\nBlock samples for manual inspection:")
        shown = 0
        for doc in docs:
            for block in doc.blocks:
                if shown >= args.print_samples:
                    break
                preview = dict(block)
                preview["text"] = preview["text"][:200]
                preview["doc_title"] = doc.doc_title
                print(json.dumps(preview, ensure_ascii=False))
                shown += 1
            if shown >= args.print_samples:
                break


if __name__ == "__main__":
    main()
